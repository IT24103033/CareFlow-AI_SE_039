from pathlib import Path
from docx import Document
D=Document('/Users/sujana/Downloads/SE3100_Individual_Section_Component-c.docx')
paras=[p.text for p in D.paragraphs if p.text.strip()]
a=max(i for i,t in enumerate(paras) if t=='Individual AI Reflection');b=paras.index('Signed Declaration')
reflection=paras[a+1:b];declaration=paras[b+1]
log=next(t for t in D.tables if t.rows[0].cells[0].text=='Task')
rows=[[c.text for c in r.cells] for r in log.rows[1:]]
for table in D.tables:
 for row in table.rows:
  if row.cells[0].text=='Signature':
   blip=row._tr.xpath('.//a:blip')[0];rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
   part=D.part.related_parts[rid];sig=Path('tmp/report-evidence/component-c-signature'+Path(part.partname).suffix);sig.write_bytes(part.blob)
p=Path('tmp/report-evidence/build_report.py');s=p.read_text()
block=''' if component=='C':
  page('16 Individual Report Component C')
  txt('Pallawala S R | IT24102852 | Appointments and resource scheduling')
  sub('Individual contribution')
  txt('I was responsible for appointment scheduling, doctor availability, conflict checking and the Action/Tool Agent. I also contributed to React and Flutter appointment interfaces and the tests and CI work associated with Component C.')
  sub('Backend, database and agent work')
  txt('My work covered doctor availability management, available-slot retrieval, appointment date/time validation, tentative booking and doctor approval before confirmation. Doctors, DoctorAvailabilities and Appointments hold the scheduling data. Transaction-based conflict checks were introduced to reduce concurrent double-booking risks.')
  txt('The Action/Tool Agent uses allow-listed FindAvailableSlots, CheckBookingConflict and CreateTentativeBooking tools. Injected provider abstractions avoid direct calls to protected controller endpoints. AppointmentActionResult and AppointmentActionStatus distinguish successful bookings, unavailable slots, conflicts, invalid requests and provider failures. AppointmentWorkflowRunner connects these contracts to the workflow layer.')
  sub('Client integration and technical learning')
  txt('I connected Flutter appointment screens to real doctor, availability and booking APIs and implemented the React appointment page. Challenges included dependency-injection registration, Git conflicts and B/C contract mismatches. I learned to align interfaces and authentication across components and to verify AI suggestions through builds, tests and API checks.')
  sub('Git evidence')
  table(['Commit','Contribution'],[['c2c3ed8','Appointment scheduling foundation.'],['34b4241','Doctor availability APIs.'],['a7fbcde','Appointment scheduling fixes and integration.']],[100,383])
  txt('Feature branch: feature/Resource-Scheduling. The supplied document identifies the PR by title; local history also contains merges #9 and #24 from this branch. Shared B/C integration changes must retain their original attribution.', 'SmallX')
  sub('Testing evidence and its limits')
  txt('Sandathi reports historical local results of 92/92 backend tests and 13/13 Flutter tests, plus a successful Flutter CI check. Her report also records a later B/C compilation mismatch in TriageOrchestrationHelper.cs. These are historical observations, not the current release status. This document update did not rerun the tests.', 'SmallX')
  txt('The consolidated report separately records the inspected baseline results. Its scheduling concurrency test was skipped under EF InMemory; the supplied report does not establish current passing PostgreSQL concurrency coverage. Confirm the final release run before submission.', 'SmallX')
  page('16 Individual AI Usage Log')
  txt('Pallawala S R | IT24102852')
  txt('The following entries summarize Sandathi\'s supplied log. Dates, tool/model names and entry-specific changes or verification evidence were not provided and must be completed by the student.', 'SmallX')
  table(['Task','AI assistance','Reported outcome'],LOG_ROWS,[102,190,191])
  txt('Student-reported review approach: AI suggestions were reviewed and adapted to the existing architecture. Her reflection describes checking project structure, authentication and interfaces, and verification through builds, automated tests, API testing and CI.', 'SmallX')
  page('16 Individual Reflection and Declaration')
  txt('Pallawala S R | IT24102852')
  sub('Individual AI reflection')
REFLECTION_LINES
  sub('Signed declaration supplied by the student')
DECLARATION_LINE
  txt('Student: Pallawala S R | IT24102852     Date: 05/10/2026', 'SmallX')
  txt('Signature:', 'SmallX')
  from reportlab.platypus import Image
  sig_image=Image(SIGNATURE_PATH)
  sig_image.drawHeight=42
  sig_image.drawWidth=42*sig_image.imageWidth/sig_image.imageHeight
  sig_image.hAlign='LEFT'
  story.append(sig_image)
  continue
'''
# Avoid apostrophe escaping in generated Python.
block=block.replace("Sandathi's supplied",'the supplied')
block=block.replace('LOG_ROWS',repr(rows)).replace('REFLECTION_LINES','\n'.join('  txt('+repr(t)+')' for t in reflection)).replace('DECLARATION_LINE','  txt('+repr(declaration)+", 'SmallX')").replace('SIGNATURE_PATH',repr(str(sig)))
needle='for num,component,name,sid,area,scope,db,api,ui,agent,tests,git in members:\n'
assert needle in s;s=s.replace(needle,needle+block)
s=s.replace('24 to 31','24 to 32').replace('32 to 34','33 to 35')
s=s.replace('A-C logs are still pending.','Sandathi supplied her contribution, task-based AI log, reflection and signed declaration; dates and tool/model details remain missing from her log. A-B logs are still pending.')
s=s.replace('D contribution, log and reflection supplied; A-C pending','C/D material supplied; A-B personal evidence pending')
# Keep reference entry concise to preserve the existing page budget.
s=s.replace("('E8','ComponentD.docx supplied by Dulshan (IT24102599) on 5 October 2026: contribution details, AI usage log, original reflection and declaration. Log dates/model details and signature remain for member confirmation.')", "('E8','Member documents: ComponentD.docx (IT24102599) and SE3100_Individual_Section_Component-c.docx (IT24102852). Contribution evidence, AI logs and original reflections; C includes a signed declaration. C/D log details still need confirmation.')")
p.write_text(s)
e=Path('tmp/report-evidence/build_erd.py');s=e.read_text().replace('assert len(reader.pages)==34','assert len(reader.pages)>=34');e.write_text(s)
print('Component C inserted with original reflection and supplied signature.')
