from pathlib import Path
from docx import Document
D=Document('/Users/sujana/Downloads/Individual Contribution Report.docx');paras=[p.text.strip() for p in D.paragraphs if p.text.strip()]
a=next(i for i,t in enumerate(paras) if t.startswith('4. AI Reflection'));b=paras.index('5. Signed Declaration');reflection=paras[a+1:b];decl=paras[b+1]
log=next(t for t in D.tables if t.rows[0].cells[0].text=='Tool');rows=[[c.text for c in r.cells] for r in log.rows[1:]]
p=Path('tmp/report-evidence/build_report.py');s=p.read_text()
block=''' if component=='A':
  page('14 Individual Report Component A')
  txt('Bandara MMSD | IT24103089 | Patient admissions, wards and shared user management')
  sub('Contribution statement')
  txt('My primary responsibility was Component A across ASP.NET Core, PostgreSQL, React, Flutter and domain analysis. I designed PatientProfiles, Wards and Admissions with EF Core, developed patient/admission APIs, and contributed to the transition from local PostgreSQL to the shared Neon database.')
  sub('Database and backend')
  txt('The admission workflow links a patient to a ward and updates OccupiedBeds in the same transaction. The current route is POST /api/Admissions/allocate-ward. It checks ward existence and capacity before saving. A transaction supports consistent writes; concurrent capacity enforcement still requires separate verification.')
  sub('React, Flutter and shared authentication')
  txt('I built staff-facing patient tiles, patient history search through GET /api/PatientProfiles/search, the ward-capacity display and the Admit to Ward modal. Mobile work enabled patients to view medical history and blood group through patient-profile API integration. I also contributed the role-based web login, backend user models and role verification. Mobile authentication and extended account management were shared work completed by a teammate.')
  sub('Domain analysis agent')
  txt('The Gemini-based domain agent combines symptoms with patient history and requests structured JSON containing risk and a ward recommendation. My work included provider failure handling and the migration from a local Ollama integration. The supplied report describes a high-risk fallback; fallback routing must be distinguished from a successful model assessment and reviewed by authorized staff.')
  sub('DevOps, evidence and challenges')
  txt('I contributed the backend GitHub Actions pipeline and investigated build failures caused by DTO changes, including AdmissionRequestDto. The supplied screenshots show project setup PR #1 merged and admissions PR #7 closed. These statuses are preserved as historical evidence rather than described as current passing CI.')
  txt('Challenges included aligning the React admission interaction with backend state changes, enforcing JSON response structure, configuring cloud database access and resolving schema/merge mismatches. The screenshots that follow document the supplied development evidence; current deployment and final test results are recorded separately.')
  page('14 Individual AI Usage Log')
  txt('Bandara MMSD | IT24103089')
  txt('Supplied tool/task records are reproduced below. Dates, exact model versions and task-specific verification results were not supplied and must be completed by the student.','SmallX')
  table(['Tool','Prompt or task','Student modification'],LOG_ROWS,[70,216,197])
  txt('Sanuthmi reports reviewing generated models and UI, wiring API calls, adjusting modal state and using local dotnet build to diagnose the DTO mismatch. These statements do not replace dated session records.','SmallX')
  page('14 Individual Reflection and Declaration')
  txt('Bandara MMSD | IT24103089')
  sub('Individual AI reflection')
REFLECTION
  sub('Signed declaration supplied by the student')
DECLARATION
  txt('Date: 06/10/2026','SmallX')
  from reportlab.platypus import Image
  signature=Image('tmp/report-evidence/a-evidence-13.png',width=150,height=48)
  signature.hAlign='LEFT';story.append(signature)
  def a_image(n,caption,width=483,maxheight=480):
   im=Image(f'tmp/report-evidence/a-evidence-{n}.png')
   iw,ih=im.imageWidth,im.imageHeight
   scale=min(width/iw,maxheight/ih)
   im.drawWidth=iw*scale;im.drawHeight=ih*scale
   story.append(im);story.append(Spacer(1,8));txt(caption,'SmallX');story.append(Spacer(1,10))
  page('14 Component A Evidence Project Setup')
  a_image(1,'Figure A1. Supplied pull-request list for Component A, showing admissions and project setup work.')
  a_image(2,'Figure A2. Project setup and CI pipeline PR #1 shown as merged. Historical development evidence.',maxheight=410)
  page('14 Component A Evidence Admissions Review')
  a_image(3,'Figure A3. Admissions PR #7 shown as closed. The discussion describes continued work through another branch; closed must not be interpreted as directly merged.')
  page('14 Component A Evidence Database')
  a_image(4,'Figure A4. Supplied pgAdmin cloud connection view; source password is masked. Connection visibility alone does not prove all cloud workflows.',width=420)
  a_image(5,'Figure A5. Supplied PatientProfiles table view showing profile and history fields.',width=420)
  a_image(6,'Figure A6. Supplied Wards table view showing capacity and occupied-bed values.',width=420)
  page('14 Component A Evidence Staff Interface')
  a_image(7,'Figure A7. Patient registration modal in the staff interface.')
  a_image(8,'Figure A8. Ward-capacity dashboard with occupancy indicators.')
  page('14 Component A Evidence Analysis and Mobile')
  left=Image('tmp/report-evidence/a-evidence-9.png',width=218,height=364)
  right=Image('tmp/report-evidence/a-evidence-10.png',width=162,height=364)
  story.append(Table([[left,right]],colWidths=[241,242]));story.append(Spacer(1,12))
  txt('Figure A9 (left). Supplied risk-analysis result with risk, ward recommendation, factors and patient history. Figure A10 (right). Flutter medical-profile screen displaying history, blood group and date of birth. These are historical UI captures.','SmallX')
  txt('The two supplied Swagger authentication screenshots were excluded because they display a token response. Provide redacted API/test evidence for final inclusion.','SmallX')
  continue
'''
block=block.replace('LOG_ROWS',repr(rows)).replace('REFLECTION','\n'.join('  txt('+repr(t)+')' for t in reflection)).replace('DECLARATION','  txt('+repr(decl)+", 'SmallX')")
needle='for num,component,name,sid,area,scope,db,api,ui,agent,tests,git in members:\n';assert needle in s;s=s.replace(needle,needle+block)
s=s.replace('24 to 36','24 to 42').replace('37 to 39','43 to 45').replace('A-B logs are still pending.','Sanuthmi supplied a tool/task log, original reflection and signed declaration; dates and exact models remain incomplete. Component B personal evidence is still pending.').replace('C/D material supplied; A-B personal evidence pending','A/C/D material supplied; B personal evidence pending')
s=s.replace('Member documents: ComponentD.docx (IT24102599) and SE3100_Individual_Section_Component-c.docx (IT24102852). Contribution evidence, AI logs and original reflections; C includes a signed declaration. C/D log details still need confirmation.','Member documents: Individual Contribution Report.docx (A), ComponentD.docx (D) and SE3100_Individual_Section_Component-c.docx (C). Contributions, logs and original reflections; A/C include signatures. Log details remain incomplete.')
p.write_text(s)
